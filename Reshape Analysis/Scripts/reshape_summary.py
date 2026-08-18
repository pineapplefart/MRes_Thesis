# -*- coding: utf-8 -*-
"""
Per-strain growth summaries for the RESHAPE analysis.

Reduces each plate's growth curve to one number — the radial extension rate Kr —
and builds a strain x substrate comparison from it.

Kr is the slope of a straight-line fit to colony radius against time, taken over
the linear phase only: the frames where the radius lies between KR_LO and KR_HI of
that plate's final radius. That deliberately excludes

  * the lag, which is a separate quantity and would drag the slope down, and
  * any plateau where the colony has run into the dish wall and stopped.

The same fit yields the lag for free (its x-intercept) and an R2 so a bad fit is
visible rather than silent.

Imperial College London
@Author: Wilson Porteus
@Date: July 2026
"""

from __future__ import annotations

from pathlib import Path

import numpy as np
import pandas as pd

from reshape_platemap import ROLE_MINUS, ROLE_SAMPLE

# Linear phase: fraction of each plate's final radius used for the fit.
KR_LO, KR_HI = 0.2, 0.8
KR_MIN_POINTS = 4          # fewer than this and the slope is not meaningful
KR_MIN_RADIUS_MM = 2.0     # below this the plate never really grew

# A substrate whose minus control grows nearly as well as its samples has not
# discriminated: the carbon (or nitrogen) source under test made little difference.
MINUS_FAIL_RATIO = 0.90

# --- exponential phase -------------------------------------------------------
# Radial extension in filamentous fungi is classically LINEAR once the peripheral
# growth zone is established; the exponential phase is the earlier expansion before
# that. Kr and mu therefore describe different parts of the same curve and are not
# interchangeable — both are reported so the choice stays visible.
#
# mu is the slope of ln(radius) against time (per hour), taken from the steepest
# well-fitting window: slide a fixed-width window along the log curve and keep the
# largest slope, the standard maximum specific growth rate.
EXP_WINDOW_H = 6.0          # width of the sliding window
EXP_MIN_RADIUS_MM = 1.0     # ignore radii below this: ln() amplifies noise near zero
                            # and would report an enormous rate for a speck appearing
EXP_MIN_R2 = 0.90           # a window must fit this well to be considered
EXP_MIN_POINTS = 4


def radial_extension_rate(times, radii) -> dict:
    """
    Fit the linear phase of one plate's growth curve.

    Returns Kr (mm/h), the lag (h, x-intercept), R2, the number of points used, and
    the final radius. Kr is NaN when the plate never grew or the phase is too short
    to fit — NaN rather than 0, so "no growth" and "not measurable" stay distinct.
    """
    t = np.asarray(times, dtype=float)
    r = np.asarray(radii, dtype=float)
    order = np.argsort(t)
    t, r = t[order], r[order]

    out = {"Kr_mm_per_h": np.nan, "Lag_h": np.nan, "Kr_R2": np.nan,
           "Kr_n_points": 0, "Final_radius_mm": float(r[-1]) if r.size else np.nan}

    if r.size < KR_MIN_POINTS or not np.isfinite(r).any():
        return out

    final = float(np.nanmax(r))
    out["Final_radius_mm"] = final
    if final < KR_MIN_RADIUS_MM:
        return out

    lo, hi = KR_LO * final, KR_HI * final
    sel = (r >= lo) & (r <= hi)
    if sel.sum() < KR_MIN_POINTS:
        # Very fast plates can jump through the window; fall back to everything
        # between the first growth and the plateau.
        sel = (r > 0) & (r <= hi)
    if sel.sum() < KR_MIN_POINTS:
        return out

    ts, rs = t[sel], r[sel]
    slope, intercept = np.polyfit(ts, rs, 1)
    pred = slope * ts + intercept
    ss_res = float(np.sum((rs - pred) ** 2))
    ss_tot = float(np.sum((rs - np.mean(rs)) ** 2))

    out["Kr_mm_per_h"] = float(slope)
    out["Kr_R2"] = 1.0 - ss_res / ss_tot if ss_tot > 0 else np.nan
    out["Kr_n_points"] = int(sel.sum())
    out["Lag_h"] = float(-intercept / slope) if slope > 0 else np.nan
    return out


def exponential_growth_rate(times, radii) -> dict:
    """
    Maximum specific growth rate mu (per hour) from ln(radius) against time.

    A fixed-width window is slid along the log-transformed curve and the steepest
    window that fits well is kept. Radii below EXP_MIN_RADIUS_MM are excluded: the
    logarithm amplifies noise near zero, so the first speck to appear would otherwise
    always look like the fastest growth on the plate.

    Returns mu, its window start, R2, and the implied doubling time.
    """
    t = np.asarray(times, dtype=float)
    r = np.asarray(radii, dtype=float)
    order = np.argsort(t)
    t, r = t[order], r[order]

    out = {"Mu_per_h": np.nan, "Mu_R2": np.nan, "Mu_window_start_h": np.nan,
           "Doubling_time_h": np.nan, "Mu_n_points": 0}

    usable = np.isfinite(r) & (r >= EXP_MIN_RADIUS_MM)
    if usable.sum() < EXP_MIN_POINTS:
        return out
    t, r = t[usable], r[usable]
    y = np.log(r)

    best = None
    for i in range(len(t)):
        j = np.searchsorted(t, t[i] + EXP_WINDOW_H, side="right")
        if j - i < EXP_MIN_POINTS:
            continue
        ts, ys = t[i:j], y[i:j]
        if np.ptp(ts) <= 0:
            continue
        slope, intercept = np.polyfit(ts, ys, 1)
        if slope <= 0:
            continue
        pred = slope * ts + intercept
        ss_tot = float(np.sum((ys - np.mean(ys)) ** 2))
        if ss_tot <= 0:
            continue
        r2 = 1.0 - float(np.sum((ys - pred) ** 2)) / ss_tot
        if r2 < EXP_MIN_R2:
            continue
        if best is None or slope > best[0]:
            best = (float(slope), float(r2), float(ts[0]), int(j - i))

    if best is None:
        return out
    slope, r2, start, n = best
    out.update({"Mu_per_h": slope, "Mu_R2": r2, "Mu_window_start_h": start,
                "Doubling_time_h": float(np.log(2) / slope) if slope > 0 else np.nan,
                "Mu_n_points": n})
    return out


def per_plate_rates(raw: pd.DataFrame) -> pd.DataFrame:
    """One row per plate: Kr, lag, fit quality and the metadata needed to group."""
    records = []
    keys = ["Condition", "Experiment", "Bio_Replicate", "Plate"]
    for key, g in raw.groupby(keys, dropna=False):
        meta = g.iloc[0]
        fit = radial_extension_rate(g["Time_h"], g["Radius_mm"])
        fit.update(exponential_growth_rate(g["Time_h"], g["Radius_mm"]))
        records.append({
            "Condition": key[0], "Experiment": key[1],
            "Bio_Replicate": key[2], "Plate": key[3],
            "Role": meta["Role"], "Strain": meta.get("Strain"),
            "Replicate": meta.get("Replicate"),
            "Carbon": meta.get("Carbon"), "Nitrogen": meta.get("Nitrogen"),
            "Minus_Control_Omits": meta.get("Minus_Control_Omits"),
            "Final_area_mm2": float(g.sort_values("Time_h")["Area_mm2"].iloc[-1]),
            "Duration_h": float(g["Time_h"].max()),
            **fit,
        })
    return pd.DataFrame(records)


def minus_control_check(rates: pd.DataFrame) -> pd.DataFrame:
    """
    Per condition, how the minus control compares with the samples.

    Ratio near or above 1 means the plate lacking the variable component grew about
    as well as the plates containing it, so that substrate did not discriminate.
    """
    rows = []
    for condition, g in rates.groupby("Condition", dropna=False):
        samples = g[g["Role"] == ROLE_SAMPLE]["Kr_mm_per_h"].dropna()
        minus = g[g["Role"] == ROLE_MINUS]["Kr_mm_per_h"].dropna()
        omits = g["Minus_Control_Omits"].dropna()
        sample_mu = g[g["Role"] == ROLE_SAMPLE]["Mu_per_h"].dropna()
        minus_mu = g[g["Role"] == ROLE_MINUS]["Mu_per_h"].dropna()

        sample_kr = float(np.median(samples)) if len(samples) else np.nan
        minus_kr = float(np.median(minus)) if len(minus) else np.nan
        ratio = minus_kr / sample_kr if sample_kr and sample_kr > 0 else np.nan

        s_mu = float(np.median(sample_mu)) if len(sample_mu) else np.nan
        m_mu = float(np.median(minus_mu)) if len(minus_mu) else np.nan
        mu_ratio = m_mu / s_mu if s_mu and s_mu > 0 else np.nan

        rows.append({
            "Condition": condition,
            "Omitted_component": omits.iloc[0] if len(omits) else "",
            "Sample_Kr_median": sample_kr,
            "Minus_Kr": minus_kr,
            "n_minus_plates": int(len(minus)),
            "Minus_over_sample": ratio,
            "Minus_control_failed": bool(ratio >= MINUS_FAIL_RATIO) if ratio == ratio else False,
            "Sample_Mu_median": s_mu,
            "Minus_Mu": m_mu,
            "Minus_over_sample_Mu": mu_ratio,
        })
    return pd.DataFrame(rows)


def strain_substrate_summary(rates: pd.DataFrame, minus: pd.DataFrame) -> pd.DataFrame:
    """
    Kr per strain per substrate, averaged the nested way: technical replicates within
    a run first, then across biological replicates, so n is the number of runs.
    """
    samples = rates[(rates["Role"] == ROLE_SAMPLE) & rates["Kr_mm_per_h"].notna()]
    if samples.empty:
        return pd.DataFrame()

    per_run = (samples.groupby(["Strain", "Condition", "Bio_Replicate"], dropna=False)
               .agg(Kr_run=("Kr_mm_per_h", "mean"),
                    Lag_run=("Lag_h", "mean"),
                    R2_run=("Kr_R2", "mean"),
                    Mu_run=("Mu_per_h", "mean"),
                    MuR2_run=("Mu_R2", "mean"),
                    Dbl_run=("Doubling_time_h", "mean"),
                    n_plates=("Plate", "count"))
               .reset_index())

    out = (per_run.groupby(["Strain", "Condition"], dropna=False)
           .agg(Kr_mm_per_h=("Kr_run", "mean"),
                Kr_SD=("Kr_run", "std"),
                Lag_h=("Lag_run", "mean"),
                R2=("R2_run", "mean"),
                Mu_per_h=("Mu_run", "mean"),
                Mu_SD=("Mu_run", "std"),
                Mu_R2=("MuR2_run", "mean"),
                Doubling_time_h=("Dbl_run", "mean"),
                n_runs=("Bio_Replicate", "nunique"),
                n_plates=("n_plates", "sum"))
           .reset_index())

    flags = minus.set_index("Condition")
    out["Minus_over_sample"] = out["Condition"].map(flags["Minus_over_sample"])
    out["Minus_control_failed"] = out["Condition"].map(
        flags["Minus_control_failed"]).fillna(False)
    out["Omitted_component"] = out["Condition"].map(flags["Omitted_component"])
    out["Exclude"] = np.where(out["Minus_control_failed"], "review", "")
    return out.sort_values(["Strain", "Kr_mm_per_h"], ascending=[True, False])


def kr_matrix(summary: pd.DataFrame, value="Kr_mm_per_h") -> pd.DataFrame:
    """Strain x substrate matrix of the chosen column."""
    if summary.empty:
        return pd.DataFrame()
    return summary.pivot_table(index="Strain", columns="Condition", values=value)


# ============================================================ output writers
FONT_NAME = "Arial"
_HDR_FILL = "1F3864"
_FLAG_FILL = "FFC7CE"      # red: minus control did not discriminate
_OK_FILL = "E2EFDA"        # green
_WARN_FILL = "FFF2CC"      # amber: few replicates


def write_summary_workbook(summary: pd.DataFrame, rates: pd.DataFrame,
                           minus: pd.DataFrame, path):
    """
    Excel summary: a strain x substrate matrix, the flagged conditions, and one
    sheet per strain. Cells are coloured so the substrates to review stand out.
    """
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
    from openpyxl.utils import get_column_letter

    thin = Side(style="thin", color="BFBFBF")
    border = Border(left=thin, right=thin, top=thin, bottom=thin)

    def header(ws, values, row=1):
        for i, h in enumerate(values, start=1):
            c = ws.cell(row=row, column=i, value=h)
            c.font = Font(name=FONT_NAME, bold=True, color="FFFFFF", size=11)
            c.fill = PatternFill("solid", fgColor=_HDR_FILL)
            c.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
            c.border = border
        ws.row_dimensions[row].height = 30

    wb = Workbook()
    wb.remove(wb.active)

    # ---- README ------------------------------------------------------
    ws = wb.create_sheet("README")
    lines = [
        "RESHAPE strain summary",
        "",
        "Kr (mm/h) is the radial extension rate: the slope of colony radius against time,",
        f"fitted over the linear phase only ({int(KR_LO*100)}-{int(KR_HI*100)}% of each plate's final radius).",
        "The lag and any plateau where the colony reached the dish wall are excluded, so Kr",
        "compares how FAST a strain extends, independent of how long it took to start.",
        "",
        "Averaging is nested: the plates of a strain within one run are technical replicates and",
        "are averaged first; SD is then taken across biological replicates, so n_runs is the",
        "number of independent runs. With n_runs = 1 the SD is blank, not zero.",
        "",
        "COLOUR CODE",
        "  red    Minus control failed: the plate without the variable component grew nearly as",
        f"         well as the samples (Kr ratio >= {MINUS_FAIL_RATIO:.2f}). That substrate did not",
        "         discriminate, so its Kr says little about the carbon or nitrogen source under",
        "         test. These are the candidates to exclude.",
        "  amber  Only one biological replicate, so there is no SD.",
        "  green  Two or more runs and the minus control behaved as expected.",
        "",
        "Kr is NaN where a plate never grew (final radius below the minimum) or the linear phase",
        "was too short to fit. NaN and 0 mean different things and are kept distinct.",
        "",
        "R2 is the fit quality for Kr. Values well below ~0.95 mean the growth curve was not",
        "straight over the fitted window and the rate should be treated with caution.",
    ]
    for i, text in enumerate(lines, start=1):
        c = ws.cell(row=i, column=1, value=text)
        c.font = Font(name=FONT_NAME, size=12 if i == 1 else 10,
                      bold=(i == 1 or text.isupper()))
    ws.column_dimensions["A"].width = 100

    # ---- matrix ------------------------------------------------------
    matrix = kr_matrix(summary)
    if not matrix.empty:
        ws = wb.create_sheet("Kr matrix")
        ws.cell(row=1, column=1, value="Radial extension rate Kr (mm/h), strain x substrate")\
            .font = Font(name=FONT_NAME, bold=True, size=13)
        header(ws, ["Strain"] + list(matrix.columns), row=3)
        flagged = set(summary.loc[summary["Minus_control_failed"] == True,
                                  ["Strain", "Condition"]].itertuples(index=False, name=None))
        for r, (strain, row) in enumerate(matrix.iterrows(), start=4):
            c = ws.cell(row=r, column=1, value=strain)
            c.font = Font(name=FONT_NAME, bold=True, size=11); c.border = border
            for j, cond in enumerate(matrix.columns, start=2):
                v = row[cond]
                cell = ws.cell(row=r, column=j,
                               value=None if pd.isna(v) else round(float(v), 3))
                cell.font = Font(name=FONT_NAME, size=11)
                cell.border = border
                cell.alignment = Alignment(horizontal="center")
                if (strain, cond) in flagged:
                    cell.fill = PatternFill("solid", fgColor=_FLAG_FILL)
        ws.column_dimensions["A"].width = 14
        for j in range(2, len(matrix.columns) + 2):
            ws.column_dimensions[get_column_letter(j)].width = 26
        ws.freeze_panes = "B4"

    # ---- flagged conditions -------------------------------------------
    ws = wb.create_sheet("Substrates to review")
    ws.cell(row=1, column=1,
            value="Conditions whose minus control grew nearly as well as the samples")\
        .font = Font(name=FONT_NAME, bold=True, size=13)
    ws.cell(row=2, column=1,
            value="The variable component made little difference here, so these substrates "
                  "are the candidates to exclude.").font = Font(name=FONT_NAME, italic=True,
                                                                size=9, color="808080")
    cols = ["Condition", "Omitted_component", "Sample_Kr_median", "Minus_Kr",
            "Minus_over_sample", "n_minus_plates", "Minus_control_failed"]
    header(ws, cols, row=3)
    ordered = minus.sort_values("Minus_over_sample", ascending=False)
    for r, (_, row) in enumerate(ordered.iterrows(), start=4):
        for j, col in enumerate(cols, start=1):
            v = row[col]
            if isinstance(v, (float, np.floating)) and pd.notna(v):
                v = round(float(v), 3)
            cell = ws.cell(row=r, column=j, value=None if pd.isna(v) else v)
            cell.font = Font(name=FONT_NAME, size=11); cell.border = border
            if row["Minus_control_failed"]:
                cell.fill = PatternFill("solid", fgColor=_FLAG_FILL)
    for j, w in enumerate([34, 20, 18, 14, 18, 16, 20], start=1):
        ws.column_dimensions[get_column_letter(j)].width = w
    ws.freeze_panes = "A4"

    # ---- one sheet per strain -----------------------------------------
    cols = ["Condition", "Omitted_component", "Kr_mm_per_h", "Kr_SD", "Lag_h", "R2",
            "n_runs", "n_plates", "Minus_over_sample", "Exclude"]
    for strain, g in summary.groupby("Strain", dropna=False):
        title = str(strain)[:28] or "unnamed"
        ws = wb.create_sheet(title)
        ws.cell(row=1, column=1, value=f"Strain {strain} — substrates ranked by Kr")\
            .font = Font(name=FONT_NAME, bold=True, size=13)
        header(ws, cols, row=3)
        for r, (_, row) in enumerate(g.iterrows(), start=4):
            for j, col in enumerate(cols, start=1):
                v = row.get(col)
                if isinstance(v, (float, np.floating)) and pd.notna(v):
                    v = round(float(v), 3)
                cell = ws.cell(row=r, column=j, value=None if pd.isna(v) else v)
                cell.font = Font(name=FONT_NAME, size=11); cell.border = border
                if col != "Condition":
                    cell.alignment = Alignment(horizontal="center")
            fill = (_FLAG_FILL if row["Minus_control_failed"]
                    else _WARN_FILL if row["n_runs"] < 2 else _OK_FILL)
            ws.cell(row=r, column=1).fill = PatternFill("solid", fgColor=fill)
        for j, w in enumerate([34, 20, 14, 12, 12, 10, 10, 10, 18, 12], start=1):
            ws.column_dimensions[get_column_letter(j)].width = w
        ws.freeze_panes = "A4"

    wb.save(path)
    return path


def plot_minus_controls(minus: pd.DataFrame, rates: pd.DataFrame, path):
    """
    Minus controls compared with one another.

    Left: their absolute rates side by side against each condition's own sample
    median, so the controls can be ranked against each other.
    Right: each as a fraction of its own samples. That is the quantity that says
    whether omitting the component mattered — a bar reaching 1.0 means the plate
    without the carbon or nitrogen source grew just as fast as the plates with it.
    """
    import matplotlib.pyplot as plt

    d = minus.dropna(subset=["Minus_Kr"]).copy()
    if d.empty:
        return None
    d = d.sort_values("Minus_over_sample", ascending=True)
    labels = [f"{c}\n(no {o})" if o else str(c)
              for c, o in zip(d["Condition"], d["Omitted_component"])]
    y = np.arange(len(d))

    fig, (ax1, ax2) = plt.subplots(
        1, 2, figsize=(15, 1.6 + 0.62 * len(d)),
        gridspec_kw={"width_ratios": [1.15, 1]})

    ax1.barh(y - 0.19, d["Minus_Kr"], height=0.36, color="#C0504D",
             label="minus control")
    ax1.barh(y + 0.19, d["Sample_Kr_median"], height=0.36, color="#9BBB59",
             label="samples (median)")
    for i, (mk, sk) in enumerate(zip(d["Minus_Kr"], d["Sample_Kr_median"])):
        if pd.notna(mk):
            ax1.text(mk, i - 0.19, f" {mk:.2f}", va="center", fontsize=8)
        if pd.notna(sk):
            ax1.text(sk, i + 0.19, f" {sk:.2f}", va="center", fontsize=8)
    ax1.set_yticks(y)
    ax1.set_yticklabels(labels, fontsize=8)
    ax1.set_xlabel("Radial extension rate Kr (mm/h)")
    ax1.set_title("Minus controls vs their own samples")
    ax1.legend(fontsize=8, loc="lower right")
    ax1.grid(axis="x", alpha=0.3)

    colours = ["#C0504D" if f else "#4F81BD" for f in d["Minus_control_failed"]]
    ax2.barh(y, d["Minus_over_sample"], color=colours, height=0.6)
    ax2.axvline(1.0, color="0.3", linewidth=1)
    ax2.axvline(MINUS_FAIL_RATIO, color="#C0504D", linestyle="--", linewidth=1.2,
                label=f"flag threshold ({MINUS_FAIL_RATIO:g})")
    for i, v in enumerate(d["Minus_over_sample"]):
        if pd.notna(v):
            ax2.text(v, i, f" {v:.2f}", va="center", fontsize=8)
    ax2.set_yticks(y)
    ax2.set_yticklabels([])
    ax2.set_xlabel("minus control Kr / sample Kr")
    ax2.set_title("Did omitting the component matter?\n1.0 = no difference at all")
    ax2.legend(fontsize=8, loc="lower right")
    ax2.grid(axis="x", alpha=0.3)

    fig.suptitle("Minus-variable controls compared", fontsize=13)
    fig.tight_layout()
    fig.savefig(path, dpi=300)
    plt.close(fig)
    return path


def plot_kr_vs_mu(rates: pd.DataFrame, path):
    """
    Linear Kr against exponential mu, one point per sample plate.

    The two describe different parts of the curve, so this shows how far they
    actually agree on your data rather than assuming they rank strains the same way.
    """
    import matplotlib.pyplot as plt

    d = rates[rates["Role"] == ROLE_SAMPLE].dropna(subset=["Kr_mm_per_h", "Mu_per_h"])
    if d.empty:
        return None

    fig, ax = plt.subplots(figsize=(8, 6.5))
    for strain, g in d.groupby("Strain", dropna=False):
        ax.scatter(g["Kr_mm_per_h"], g["Mu_per_h"], s=28, alpha=0.75, label=str(strain))
    if len(d) > 2:
        r = float(np.corrcoef(d["Kr_mm_per_h"], d["Mu_per_h"])[0, 1])
        ax.set_title(f"Linear vs exponential rate, per plate   (Pearson r = {r:.2f})")
    else:
        ax.set_title("Linear vs exponential rate, per plate")
    ax.set_xlabel("Kr — linear phase (mm/h)")
    ax.set_ylabel("mu — exponential phase (per h)")
    ax.grid(alpha=0.3)
    ax.legend(fontsize=9, title="Strain")
    fig.tight_layout()
    fig.savefig(path, dpi=300)
    plt.close(fig)
    return path


def plot_kr_heatmap(summary: pd.DataFrame, path, labels: dict | None = None,
                    failed: bool | None = None, subtitle: str | None = None):
    """
    Strain x substrate heatmap of Kr.

    `failed` selects which substrates to include, by their minus-control result:
      None   every substrate
      False  only those whose minus control behaved as expected — the usable set
      True   only those whose minus control grew nearly as well as the samples

    Splitting them keeps the substrates you can interpret from being read on the
    same colour scale as the ones you cannot. Returns None, writing nothing, when
    no substrate falls in the requested group.
    """
    import matplotlib.pyplot as plt

    if failed is not None:
        summary = summary[summary["Minus_control_failed"].fillna(False) == failed]
    if summary.empty:
        return None

    matrix = kr_matrix(summary)
    if matrix.empty:
        return None
    # Only mark individual cells when both kinds are present on the one figure.
    flagged = (set() if failed is not None else
               set(summary.loc[summary["Minus_control_failed"] == True,
                               ["Strain", "Condition"]].itertuples(index=False, name=None)))

    fig, ax = plt.subplots(figsize=(1.6 * len(matrix.columns) + 4, 0.8 * len(matrix) + 3))
    data = matrix.values.astype(float)
    im = ax.imshow(data, cmap="viridis", aspect="auto")

    ax.set_xticks(range(len(matrix.columns)))
    shown = [ (labels or {}).get(c, c) for c in matrix.columns ]
    ax.set_xticklabels(shown, rotation=35, ha="right", fontsize=10)
    ax.set_yticks(range(len(matrix.index)))
    ax.set_yticklabels(matrix.index, fontsize=10)

    finite = data[np.isfinite(data)]
    mid = (finite.max() + finite.min()) / 2 if finite.size else 0
    for i, strain in enumerate(matrix.index):
        for j, cond in enumerate(matrix.columns):
            v = data[i, j]
            if not np.isfinite(v):
                ax.text(j, i, "—", ha="center", va="center", color="0.6", fontsize=10)
                continue
            marker = "*" if (strain, cond) in flagged else ""
            ax.text(j, i, f"{v:.2f}{marker}", ha="center", va="center", fontsize=9,
                    color="white" if v < mid else "black")

    if subtitle is None:
        subtitle = {False: "substrates whose minus control behaved as expected",
                    True: "substrates whose minus control grew nearly as well as the "
                          "samples — interpret with caution"}.get(
                        failed, "* minus control did not discriminate")
    ax.set_title(f"Radial extension rate Kr (mm/h)\n{subtitle}", fontsize=12)
    fig.colorbar(im, ax=ax, label="Kr (mm/h)", shrink=0.8)
    fig.tight_layout()
    Path(path).parent.mkdir(parents=True, exist_ok=True)
    fig.savefig(path, dpi=300)
    plt.close(fig)
    return path
